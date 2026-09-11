"""Format-agnostic, hierarchy-aware BOQ line-item extractor.

Confirmed working against five real, structurally different tendered-BOQ
files (two 2-row-split-header layouts, a 7-column single-row layout, an
11-column layout using roman-numeral top-level sections plus revision
columns, and a 10-column layout with Phase-1/Phase-2/Total splits) by
locating each column's PURPOSE (item no / description / uom / quantity /
rate / amount) via header-TEXT matching instead of a fixed column
position -- see detect_columns() below.

This module answers a different question than app/services/boq_extractor.py
(Phase 1's extractor): boq_extractor.py sums CATEGORY TOTALS across a whole
sheet by toggling on all-caps section-header rows ("REINFORCED CEMENT
CONCRETE" -> everything below is 'rcc' until the next header). That's
efficient for the ~7 categories it targets, but section headers only tell
you what SECTION a row is in, not what's actually in a bare row like
"9th floor" a few lines under it -- and it doesn't extract every line
item as an individually classifiable, individually carbon-costed record.
This module keeps the running section/category/subsection CONTEXT for
every row (not just a boolean "am I still under a header"), and returns
one record per genuine billable line item with its own context attached
-- so a downstream classifier (classifier.py) sees "REINFORCED CEMENT
CONCRETE | SDC / Free flow concrete... | 9th floor" instead of just
"9th floor", which is what actually makes bare context-dependent rows
classifiable at all.

Performance note: reads the whole sheet ONCE via
ws.iter_rows(values_only=True) into a plain list of tuples, then works
off that list by plain index -- same pattern boq_extractor.py already
uses. An earlier version of this module used ws.cell(row, column).value
random-access on a read_only worksheet for header detection AND the main
extraction loop; that pattern is known to be extremely slow on openpyxl's
read_only sheets (each .cell() call re-parses from the underlying XML
stream rather than being O(1)) and timed out on a real ~2,600-row BOQ.
Confirmed fixed by switching to iter_rows() + list indexing: full-sheet
extraction now runs in well under a second.

Row taxonomy (same across every real file checked):
  - SECTION header  : Item No. is a single letter (A, B, C...)          -> level 1
  - CATEGORY header : Item No. is a bare integer (1, 2, 3...)           -> level 2
  - SUBSECTION hdr  : Item No. is "N.NN" or "N.N" (e.g. 1.01, 1.1)      -> level 3
  - ROMAN header    : Item No. is a roman numeral (I, II, III...)      -> level 1/2
                       (used as a top-level section marker in some files,
                       e.g. "I CIVIL WORKS", and as a sub-item marker in
                       others -- disambiguated by whether the row itself
                       is billable, see extract_billable_items())
  - NOTE rows       : no UOM AND no qty/rate/amount                     -> skip (boilerplate text)
  - Sub-Total rows  : "Sub-Total" / "Total" / "Carried Forward" etc.    -> skip
  - PAYMENT-SPLIT   : "Supply 70% (Material supply)" etc.               -> skip (a billing
                       milestone breakdown of the parent item's own rate/amount,
                       not a distinct material or work type)
  - BILLABLE row    : has a UOM AND a real quantity or rate/amount      -> extracted

Known limitation (same one boq_extractor.py's own docstring flags for its
supplementary-sheet filter): the header-detection keywords and row-
taxonomy rules here were derived from BOQs seen so far (Botanico,
Ecopolitan, and three other real files with different layouts). A BOQ
from a company with very different conventions may need new keywords
added to _HEADER_KEYWORDS or a new row-classification case -- this is
built to be extended, not treated as a closed, complete list.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from typing import Optional

import openpyxl

_ROMAN_RE = re.compile(r"^(i|ii|iii|iv|v|vi|vii|viii|ix|x|xi|xii|xiii|xiv|xv)$", re.IGNORECASE)
_LETTER_RE = re.compile(r"^[A-Z]$")
_SUBSECTION_RE = re.compile(r"^\d+\.\d+[a-z]?$")

_PAYMENT_SPLIT_RE = re.compile(
    r"^\s*(supply|installation|materia?l\s*supply)\s*\d{1,3}\s*%\s*\(", re.IGNORECASE
)
_SUBTOTAL_RE = re.compile(
    r"^(sub[\s\-]?total|total|grand\s*total|carried\s*forward|c/f|b/f)\s*[:\-]?\s*$", re.IGNORECASE
)

_HEADER_KEYWORDS = {
    "item_no": ["item no", "item code", "sl no", "sl.no", "sr no", "s.no", "s. no"],
    "description": ["description of work", "description of items", "description"],
    "uom": ["uom", "unit of measure", "unit"],
    "qty": ["quantity", "qty"],
    "rate": ["rate"],
    "amount": ["amount", "total value", "value"],
}


@dataclass
class BoqLineItem:
    """One genuine billable BOQ row, with its full hierarchy context."""

    row: int
    item_no: Optional[str]
    section: Optional[str]
    category: Optional[str]
    subsection: Optional[str]
    local_group: Optional[str]  # e.g. "Toilet Walls" -- see extract_billable_items' section-lookahead comment
    description: str
    enriched_description: str  # section > category > subsection > local_group > own description, " | "-joined
    uom: Optional[str]
    qty: Optional[float]
    rate: Optional[float]
    amount: Optional[float]


def is_payment_split_row(desc) -> bool:
    return bool(_PAYMENT_SPLIT_RE.match(str(desc).strip()))


def _is_number(x) -> bool:
    return isinstance(x, (int, float)) and not isinstance(x, bool)


def _nonzero(x) -> bool:
    return _is_number(x) and x not in (0, 0.0)


def classify_item_no(item_no) -> Optional[str]:
    if item_no is None:
        return None
    s = str(item_no).strip()
    if _LETTER_RE.match(s):
        return "section"
    if _SUBSECTION_RE.match(s):
        return "subsection"
    if _ROMAN_RE.match(s):
        return "roman"
    try:
        f = float(s)
        if f == int(f):
            return "category"
    except ValueError:
        pass
    return None


def _cell(row: tuple, col_idx: int):
    """0-indexed access into a row tuple from iter_rows(values_only=True),
    tolerant of a row shorter than expected (ragged rows do happen)."""
    return row[col_idx] if col_idx is not None and col_idx < len(row) else None


def find_header_band(rows: list[tuple], max_scan: int = 15) -> tuple[int, int]:
    """Locates the header band by row INDEX into `rows` (0-indexed).
    Returns (anchor_idx, band_end_idx); data starts at band_end_idx + 1.
    """
    n_rows = len(rows)
    anchor = None
    for idx in range(min(max_scan, n_rows)):
        row = rows[idx]
        joined = " ".join(str(v) for v in row if v).lower()
        has_item = any(k in joined for k in _HEADER_KEYWORDS["item_no"])
        has_desc = "description" in joined
        if has_item and has_desc:
            anchor = idx
            break
    if anchor is None:
        anchor = 2  # fallback (0-indexed row 2 == the 3rd row), matches the most common observed layout

    # Extend the band forward only while the row is a genuine sub-header:
    # ALWAYS blank in columns 0-1 (item_no/description). A real section
    # header ("A" | "STRUCTURE WORKS") is short text too, but is NEVER
    # blank in those columns -- without this guard the first real section
    # marker right after the header gets swallowed into the header band
    # and silently skipped (confirmed as a real bug against one BOQ file).
    band_end = anchor
    for idx in range(anchor + 1, min(anchor + 4, n_rows)):
        row = rows[idx]
        non_empty = [v for v in row if v is not None and str(v).strip() != ""]
        if not non_empty:
            continue
        col0_blank = _cell(row, 0) in (None, "")
        col1_blank = _cell(row, 1) in (None, "")
        looks_like_subheader = col0_blank and col1_blank and all(
            isinstance(v, str) and len(v.strip()) <= 20 for v in non_empty
        )
        if looks_like_subheader:
            band_end = idx
        else:
            break
    return anchor, band_end


def _match_purpose(cell_text: str) -> Optional[str]:
    if not cell_text:
        return None
    t = str(cell_text).strip().lower()
    best_purpose, best_len = None, 0
    for purpose, keywords in _HEADER_KEYWORDS.items():
        for kw in keywords:
            if kw in t and len(kw) > best_len:
                best_purpose, best_len = purpose, len(kw)
    return best_purpose


def detect_columns(rows: list[tuple], anchor_idx: int, band_end_idx: int) -> dict[str, int]:
    """Scans every cell in the header band and assigns each column index
    (0-indexed) to a purpose. When several columns match the same purpose
    (revision or phase-split BOQs often have more than one Qty/Amount
    pair), prefers a column whose header text mentions 'total'; otherwise
    keeps the LAST (rightmost) match, since later revision/phase columns
    are typically appended to the right and represent the current figure.

    Merged-cell forward-fill (added after a real, sizeable bug against a
    real Phase-1/Phase-2-split BOQ): openpyxl's read_only
    iter_rows(values_only=True) -- used here deliberately for performance,
    see the module docstring -- returns the label text ONLY in the
    top-left cell of a merged range and None for every other cell inside
    it. A 3-row header band like:
        row1: ... | 'As per VO#02'(spans 4-9)         | 'Current Variation'(10-11) | ...
        row2: ... | 'Quantity'(spans 4-6) | 'Amount'(7-9)     | 'Quantity'(10-11)  | ...
        row3: ... | 'Ph-1' | 'Ph-2' | 'Total'          | 'Ph 1' | 'Ph 2'           | ...
    reads column 6's row2 cell as None (it's merged into column 4's
    "Quantity" cell), so column 6's OWN joined header text is just
    "Total" -- not "Quantity Total" -- and never gets recognized as a
    'qty' column at all, let alone the preferred TOTAL one. Confirmed on
    a real file (Ecopolitan_BOQ.xlsx): column detection silently fell
    back to a Phase-1-ONLY quantity column instead of the true
    Phase-1+Phase-2 total column, undercounting quantity on the majority
    of billable rows -- including already-computed, already-trusted rows
    (154 of 164 contributing RCC lines; both reinforcement-steel summary
    lines, off by ~16%) -- not just the classifier-coverage gap this was
    originally being investigated for.
    Fix: forward-fill each header-band row's blank cells with the last
    non-blank text seen so far IN THAT ROW, capped to columns that have
    real content somewhere in the header band (so fill doesn't run on
    into genuinely empty trailing columns past the real table). This
    reconstructs what a human reading the merged header visually sees --
    "Quantity" spanning columns 4-6 -- without needing openpyxl's merged-
    cell range API, which isn't available on read_only worksheets anyway.
    """
    n_cols = max((len(r) for r in rows[anchor_idx : band_end_idx + 1]), default=0)
    # Real content bound: don't forward-fill past the last column that
    # has ANY text anywhere in the header band, or a merged label at the
    # end of the real table would otherwise bleed into blank padding
    # columns openpyxl sometimes reports beyond the sheet's actual data.
    last_real_col = -1
    for idx in range(anchor_idx, band_end_idx + 1):
        row = rows[idx]
        for c in range(n_cols):
            v = _cell(row, c)
            if v is not None and str(v).strip() != "":
                last_real_col = max(last_real_col, c)

    col_text: dict[int, list[str]] = {c: [] for c in range(n_cols)}
    for idx in range(anchor_idx, band_end_idx + 1):
        row = rows[idx]
        carry: Optional[str] = None
        for c in range(n_cols):
            v = _cell(row, c)
            text = str(v).strip() if v is not None and str(v).strip() != "" else None
            if text is not None:
                carry = text
            elif carry is not None and c <= last_real_col:
                text = carry  # merged-cell fill -- see docstring above
            if text is not None:
                col_text[c].append(text)

    purpose_hits: dict[str, list[tuple[int, str]]] = {p: [] for p in _HEADER_KEYWORDS}
    for c in range(n_cols):
        joined = " ".join(col_text[c])
        purpose = _match_purpose(joined)
        if purpose:
            purpose_hits[purpose].append((c, joined))

    chosen: dict[str, int] = {}
    for purpose, hits in purpose_hits.items():
        if not hits:
            continue
        total_hits = [h for h in hits if "total" in h[1].lower()]
        chosen[purpose] = (total_hits[-1] if total_hits else hits[-1])[0]
    return chosen


def _next_real_row(rows: list[tuple], from_idx: int, col_desc: int) -> Optional[tuple]:
    """First row after from_idx that would actually reach the main
    extraction loop's classification step -- i.e. survives the same
    desc-blank / subtotal / payment-split skips extract_billable_items
    itself applies. Used only for the one-row lookahead described below;
    kept in sync with the skip conditions at the top of that loop.
    """
    for j in range(from_idx + 1, len(rows)):
        desc = _cell(rows[j], col_desc)
        if desc is None or (isinstance(desc, str) and not desc.strip()):
            continue
        if isinstance(desc, str) and _SUBTOTAL_RE.match(desc.strip()):
            continue
        if is_payment_split_row(desc):
            continue
        return rows[j]
    return None


def extract_billable_items(path: str, sheet_name: str) -> list[BoqLineItem]:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        # Explicit close -- Windows enforces file locks strictly (same
        # issue documented in boq_extractor.py's own docstring), so an
        # upload-and-process endpoint that later tries to delete its temp
        # file will fail with PermissionError unless this is closed here.
        wb.close()

    anchor_idx, band_end_idx = find_header_band(rows)
    cols = detect_columns(rows, anchor_idx, band_end_idx)

    col_item_no = cols.get("item_no", 0)
    col_desc = cols.get("description", 1)
    col_uom = cols.get("uom")
    col_qty = cols.get("qty")
    col_rate = cols.get("rate")
    col_amount = cols.get("amount")

    context_stack = {"section": None, "category": None, "subsection": None, "local_group": None}
    records: list[BoqLineItem] = []

    for idx in range(band_end_idx + 1, len(rows)):
        row = rows[idx]
        item_no = _cell(row, col_item_no)
        desc = _cell(row, col_desc)
        uom = _cell(row, col_uom)
        qty = _cell(row, col_qty)
        rate = _cell(row, col_rate)
        amount = _cell(row, col_amount)

        if desc is None or (isinstance(desc, str) and not desc.strip()):
            continue
        if isinstance(desc, str) and _SUBTOTAL_RE.match(desc.strip()):
            continue
        if is_payment_split_row(desc):
            continue

        kind = classify_item_no(item_no)

        if kind in ("section", "roman"):
            is_billable_probe = (uom is not None and str(uom).strip() != "") and (
                _nonzero(rate) or _nonzero(amount) or _nonzero(qty)
            )
            if kind == "roman" and is_billable_probe:
                pass  # sub-item, not a new section -- fall through
            elif kind == "section" and context_stack["subsection"] is not None:
                # Bare-uppercase-letter row ("A", "B"...) encountered WHILE
                # a numbered subsection is already active. Confirmed via
                # direct inspection of real Ecopolitan/Botanico BOQs: a
                # genuine top-level section marker (the same bare-letter
                # format, e.g. "A STRUCTURE WORKS") is always immediately
                # followed by a bare-integer CATEGORY row -- but a row like
                # "A Toilet Walls" or "B Kitchen dado", used to locally
                # relabel a group of floors WITHIN an already-active
                # subsection (whose own row usually carries the real
                # material spec, e.g. "3.03 Providing... Ceramic tile...
                # dadoing..."), is always followed directly by lowercase-
                # lettered per-floor billable rows instead. Checked across
                # every bare-letter row in both real files: the "next row
                # is a category" signal separates the two cases with no
                # exceptions in either file. Treating the second case as a
                # full section reset (the old behaviour) silently discarded
                # the active subsection's material spec for every floor row
                # under it -- e.g. 1,395 sqm of toilet-wall tiling on one
                # real project alone, dropped to an unclassifiable "Toilet
                # Walls | Floor -4" with no material text left at all.
                if classify_item_no(_cell(_next_real_row(rows, idx, col_desc) or (), col_item_no)) == "category":
                    context_stack = {"section": str(desc).strip(), "category": None, "subsection": None, "local_group": None}
                    continue
                context_stack["local_group"] = str(desc).strip()
                continue
            else:
                context_stack = {"section": str(desc).strip(), "category": None, "subsection": None, "local_group": None}
                continue
        elif kind == "category":
            context_stack["category"] = str(desc).strip()
            context_stack["subsection"] = None
            context_stack["local_group"] = None
            continue
        elif kind == "subsection":
            context_stack["subsection"] = str(desc).strip()
            context_stack["local_group"] = None
            # falls through -- a subsection header row can also itself be billable

        is_billable = (uom is not None and str(uom).strip() != "") and (
            _nonzero(rate) or _nonzero(amount) or _nonzero(qty)
        )
        if not is_billable:
            continue

        context_parts = [
            context_stack["section"],
            context_stack["category"],
            context_stack["subsection"],
            context_stack["local_group"],
        ]
        context_text = " | ".join(p for p in context_parts if p)
        own_text = str(desc).strip()
        enriched = f"{context_text} | {own_text}" if context_text else own_text

        records.append(
            BoqLineItem(
                row=idx + 1,  # report as 1-indexed, matching the spreadsheet's own row numbers
                item_no=str(item_no).strip() if item_no is not None else None,
                section=context_stack["section"],
                category=context_stack["category"],
                subsection=context_stack["subsection"],
                local_group=context_stack["local_group"],
                description=own_text,
                enriched_description=enriched,
                uom=str(uom).strip() if uom is not None else None,
                qty=float(qty) if _is_number(qty) else None,
                rate=float(rate) if _is_number(rate) else None,
                amount=float(amount) if _is_number(amount) else None,
            )
        )

    return records


def extract_from_first_primary_sheet(path: str, supplementary_sheet_patterns: list[str]) -> tuple[str, list[BoqLineItem]]:
    """Convenience helper: picks the first sheet whose name doesn't match
    any of the given supplementary-sheet patterns and extracts from it.
    Reuse boq_extractor.SUPPLEMENTARY_SHEET_PATTERNS as the patterns list
    if you want the same primary-sheet filtering Phase 1's extractor
    applies. Raises ValueError if every sheet looks supplementary.
    """
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        sheet_names = wb.sheetnames
    finally:
        wb.close()

    for name in sheet_names:
        if not any(re.search(pat, name, re.IGNORECASE) for pat in supplementary_sheet_patterns):
            return name, extract_billable_items(path, name)

    raise ValueError(f"No primary sheet found in {path} -- every sheet matched a supplementary pattern")