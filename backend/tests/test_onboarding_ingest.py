"""Workstream 05: tests for app/services/onboarding/ingest.py -- raw
item-code workbook parsing (flexible header detection) and
description-based deduplication.
"""

from __future__ import annotations

import openpyxl
import pytest

from app.services.onboarding.ingest import dedupe_by_description, parse_item_code_workbook


def test_dedupe_groups_identical_description_and_unit():
    rows = {
        "1001": {"desc": "RCC M30 footing", "unit": "CUM"},
        "1002": {"desc": "rcc m30   footing", "unit": "cum"},  # same text, different case/spacing/unit-case
        "1003": {"desc": "MS Angle section", "unit": "KG"},
    }
    result = dedupe_by_description(rows)
    assert len(result.deduped) == 2
    # Both 1001 and 1002 collapse under whichever code came first.
    groups_by_size = sorted(result.code_groups.values(), key=len)
    assert groups_by_size[-1] == ["1001", "1002"]
    assert groups_by_size[0] == ["1003"]


def test_dedupe_keeps_distinct_units_separate():
    # Same description text, but a different unit is a genuinely
    # different line item (e.g. priced per Cum vs. per Kg) -- must NOT
    # be collapsed together.
    rows = {
        "A": {"desc": "Structural steel", "unit": "KG"},
        "B": {"desc": "Structural steel", "unit": "MT"},
    }
    result = dedupe_by_description(rows)
    assert len(result.deduped) == 2


def test_dedupe_preserves_every_original_code():
    rows = {f"C{i}": {"desc": "Same text", "unit": "NOS"} for i in range(5)}
    result = dedupe_by_description(rows)
    assert len(result.deduped) == 1
    (group,) = result.code_groups.values()
    assert sorted(group) == [f"C{i}" for i in range(5)]


def _write_workbook(path, header_row_offset: int, rows: list[tuple]):
    wb = openpyxl.Workbook()
    ws = wb.active
    for _ in range(header_row_offset - 1):
        ws.append(["(cover page / notes row, not a header)"])
    ws.append(["Item Code", "Description", "Unit"])
    for r in rows:
        ws.append(list(r))
    wb.save(path)


def test_parse_item_code_workbook_finds_header_row_not_on_row_one(tmp_path):
    path = tmp_path / "items.xlsx"
    _write_workbook(path, header_row_offset=3, rows=[
        ("I-001", "RCC M30 footing", "CUM"),
        ("I-002", "MS Angle section", "KG"),
    ])
    rows = parse_item_code_workbook(str(path))
    assert len(rows) == 2
    assert rows["I-001"]["desc"] == "RCC M30 footing"
    assert rows["I-001"]["unit"] == "CUM"
    assert rows["I-001"]["contributes"] == "need_review"


def test_parse_item_code_workbook_auto_assigns_code_when_missing(tmp_path):
    path = tmp_path / "items_no_code.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Description", "Unit"])
    ws.append(["Some material", "KG"])
    wb.save(path)

    rows = parse_item_code_workbook(str(path))
    assert len(rows) == 1
    (code,) = rows.keys()
    assert code.startswith("AUTO-")


def test_parse_item_code_workbook_skips_blank_description_rows(tmp_path):
    path = tmp_path / "items_blank.xlsx"
    _write_workbook(path, header_row_offset=1, rows=[
        ("I-001", "Real item", "KG"),
        ("I-002", None, "KG"),
        ("I-003", "", "KG"),
    ])
    rows = parse_item_code_workbook(str(path))
    assert list(rows.keys()) == ["I-001"]


def test_parse_item_code_workbook_raises_when_no_header_found(tmp_path):
    path = tmp_path / "not_an_item_list.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    for i in range(20):
        ws.append([f"random cell {i}"])
    wb.save(path)

    with pytest.raises(ValueError):
        parse_item_code_workbook(str(path))